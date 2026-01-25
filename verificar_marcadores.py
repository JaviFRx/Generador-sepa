"""Script para verificar marcadores en la plantilla Word"""
from docx import Document
import re

plantilla = r"d:\Sepas\docs\plantilla_domiciliacion_sepa.docx"

doc = Document(plantilla)

print("=" * 80)
print("BUSCANDO MARCADORES {{}} EN LA PLANTILLA")
print("=" * 80)

# Buscar en párrafos
print("\n--- PÁRRAFOS ---")
for i, paragraph in enumerate(doc.paragraphs):
    if '{{' in paragraph.text or 'referencia' in paragraph.text.lower():
        print(f"Párrafo {i}: {paragraph.text}")
        marcadores = re.findall(r'\{\{(\w+)\}\}', paragraph.text)
        if marcadores:
            print(f"  -> Marcadores: {marcadores}")

# Buscar en tablas
print("\n--- TABLAS ---")
for t_idx, table in enumerate(doc.tables):
    for r_idx, row in enumerate(table.rows):
        for c_idx, cell in enumerate(row.cells):
            for p_idx, paragraph in enumerate(cell.paragraphs):
                if '{{' in paragraph.text or 'referencia' in paragraph.text.lower():
                    print(f"Tabla {t_idx}, Fila {r_idx}, Celda {c_idx}: {paragraph.text}")
                    marcadores = re.findall(r'\{\{(\w+)\}\}', paragraph.text)
                    if marcadores:
                        print(f"  -> Marcadores: {marcadores}")

print("\n" + "=" * 80)
