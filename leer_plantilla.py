from docx import Document

doc = Document('d:/Sepas/docs/plantilla_domiciliacion_sepa.docx')

print('=== CONTENIDO DE LA PLANTILLA ===')
print()

for i, para in enumerate(doc.paragraphs, 1):
    if para.text.strip():
        print(f'[{i}] {para.text}')

print()
print('=== MARCADORES ENCONTRADOS ===')
import re
for para in doc.paragraphs:
    marcadores = re.findall(r'\{\{(\w+)\}\}', para.text)
    if marcadores:
        print(f'Párrafo: "{para.text[:60]}..."')
        print(f'  Marcadores: {marcadores}')
        print()
