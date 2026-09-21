"""
Módulo para leer datos desde archivos Excel
"""

import openpyxl
from datetime import datetime, date, time
from pathlib import Path
from typing import List, Dict


class LectorExcel:
    """Clase para leer datos de archivos Excel"""
    
    def __init__(self, ruta_archivo: str):
        """
        Inicializa el lector de Excel
        
        Args:
            ruta_archivo: Ruta al archivo Excel
        """
        self.ruta_archivo = Path(ruta_archivo)
        self.workbook = None
        self.sheet = None
        
    def leer_datos(self) -> List[Dict]:
        """
        Lee los datos del Excel y retorna una lista de diccionarios
        
        Formato esperado del Excel:
        - Primera fila: encabezados (nombre, apellido, email, dni, fecha, etc.)
        - Filas siguientes: datos
        
        Returns:
            Lista de diccionarios con los datos de cada fila
        """
        try:
            # Abrir el archivo Excel
            self.workbook = openpyxl.load_workbook(self.ruta_archivo)
            self.sheet = self.workbook.active
            
            # Leer encabezados (primera fila)
            encabezados = []
            for cell in self.sheet[1]:
                encabezados.append(cell.value if cell.value else f"columna_{cell.column}")
            
            # Leer datos
            datos = []
            for row_idx in range(2, self.sheet.max_row + 1):
                fila = self.sheet[row_idx]
                
                # Crear diccionario para esta fila
                registro = {}
                for idx, cell in enumerate(fila):
                    if idx < len(encabezados):
                        # Convertir el valor a string si no es None
                        valor = cell.value
                        registro[encabezados[idx]] = self._formatear_valor(valor)
                
                # Solo agregar si hay al menos un dato
                if any(registro.values()):
                    datos.append(registro)
            
            self.workbook.close()
            return datos
            
        except Exception as e:
            if self.workbook:
                self.workbook.close()
            raise Exception(f"Error al leer el archivo Excel: {str(e)}")
    
    @staticmethod
    def _formatear_valor(valor) -> str:
        """
        Convierte el valor de una celda a texto legible.

        Las celdas de fecha/hora de Excel llegan como datetime y al hacer
        str() producen cadenas del tipo "2026-06-29 22:46:23.014000",
        que no deben aparecer en el documento final.
        """
        if valor is None:
            return ""
        if isinstance(valor, datetime):
            return valor.strftime("%d/%m/%Y")
        if isinstance(valor, date):
            return valor.strftime("%d/%m/%Y")
        if isinstance(valor, time):
            return valor.strftime("%H:%M")
        return str(valor)

    def obtener_columnas(self) -> List[str]:
        """
        Obtiene la lista de columnas (encabezados) del Excel
        
        Returns:
            Lista con los nombres de las columnas
        """
        try:
            self.workbook = openpyxl.load_workbook(self.ruta_archivo)
            self.sheet = self.workbook.active
            
            encabezados = []
            for cell in self.sheet[1]:
                encabezados.append(cell.value if cell.value else f"columna_{cell.column}")
            
            self.workbook.close()
            return encabezados
            
        except Exception as e:
            if self.workbook:
                self.workbook.close()
            raise Exception(f"Error al leer columnas: {str(e)}")
