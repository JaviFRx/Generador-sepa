from docx import Document

# Abrir la plantilla
doc = Document('d:/Sepas/docs/plantilla_domiciliacion_sepa.docx')

# Mapeo de textos a buscar y marcadores a reemplazar
campos_a_reemplazar = [
    {
        'buscar': 'Referencia   de   la   orden de   domiciliación:',
        'marcador': '{{nombre_alumno}}_Robotica'
    },
    {
        'buscar': 'Nombre del deudor/es / Debtor\'s name',
        'marcador': '{{nombre_deudor}}'
    },
    {
        'buscar': 'Dirección del deudor /Address of the debtor',
        'marcador': '{{direccion_deudor}}'
    },
    {
        'buscar': 'Código postal - Población - Provincia / Postal Code - City - Town',
        'marcador': '{{codigo_postal_poblacion}}'
    },
    {
        'buscar': 'País del deudor / Country of the debtor',
        'marcador': '{{pais_deudor}}'
    },
    {
        'buscar': 'Swift BIC / Swift BI C',
        'marcador': '{{swift_bic}}'
    }
]

print('=== ACTUALIZANDO PLANTILLA ===\n')

# Recorrer todos los párrafos
for para in doc.paragraphs:
    texto_original = para.text
    
    # Buscar cada campo
    for campo in campos_a_reemplazar:
        if campo['buscar'] in texto_original:
            # Reemplazar las líneas de guiones por el marcador
            nuevo_texto = texto_original.split('\n')[0]  # Tomar solo la primera línea
            nuevo_texto = nuevo_texto + ' ' + campo['marcador']
            
            # Limpiar el párrafo
            for run in para.runs:
                run.text = ""
            
            # Agregar el nuevo texto
            if para.runs:
                para.runs[0].text = nuevo_texto
            else:
                para.add_run(nuevo_texto)
            
            print(f'✓ Campo: {campo["buscar"][:40]}...')
            print(f'  Marcador: {campo["marcador"]}')
            print()
            break

# Guardar
doc.save('d:/Sepas/docs/plantilla_actualizada.docx')
print('\n✓ Plantilla guardada como: plantilla_actualizada.docx')
print('  Revísala y luego renómbrala a: plantilla_domiciliacion_sepa.docx')
